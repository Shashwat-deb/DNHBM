"""Central server: REST API + SQLite + dashboard host (FR-01, 05-13)."""
import csv
import hmac
import io
import json
import logging
import os
import sqlite3
import time
from datetime import datetime, timezone

from flask import Flask, Response, g, jsonify, request, send_from_directory

HERE = os.path.dirname(os.path.abspath(__file__))
METRIC_FIELDS = ["cpu_usage", "memory_usage", "network_throughput",
                 "packet_loss", "connection_count"]
# Accepted aliases from the test data in Lab-5 (cpu, mem, bandwidth, ...)
ALIASES = {"cpu": "cpu_usage", "mem": "memory_usage",
           "bandwidth": "network_throughput", "connections": "connection_count"}
RANGES = {"cpu_usage": (0, 100), "memory_usage": (0, 100), "packet_loss": (0, 100),
          "network_throughput": (0, None), "connection_count": (0, None)}
DEFAULT_INTERVAL = 5
FRESH_INTERVALS = 3  # node is offline when last push older than 3 intervals

logger = logging.getLogger("server")

SCHEMA = """
CREATE TABLE IF NOT EXISTS nodes (
    node_id TEXT PRIMARY KEY, hostname TEXT, os TEXT,
    interval_s REAL DEFAULT 5, registered_at TEXT, last_seen TEXT);
CREATE TABLE IF NOT EXISTS metrics (
    metric_id INTEGER PRIMARY KEY AUTOINCREMENT,
    node_id TEXT NOT NULL REFERENCES nodes(node_id), timestamp TEXT NOT NULL,
    cpu_usage REAL, memory_usage REAL, network_throughput REAL,
    packet_loss REAL, connection_count INTEGER);
CREATE INDEX IF NOT EXISTS idx_metrics_node_ts ON metrics(node_id, timestamp);
CREATE TABLE IF NOT EXISTS thresholds (
    threshold_id INTEGER PRIMARY KEY AUTOINCREMENT,
    metric_type TEXT NOT NULL UNIQUE, limit_value REAL NOT NULL);
CREATE TABLE IF NOT EXISTS alerts (
    alert_id INTEGER PRIMARY KEY AUTOINCREMENT,
    node_id TEXT NOT NULL REFERENCES nodes(node_id),
    threshold_id INTEGER NOT NULL REFERENCES thresholds(threshold_id),
    message TEXT, timestamp TEXT NOT NULL, is_active INTEGER DEFAULT 1);
"""


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


def setup_server_logging(log_path=None):
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
        sh = logging.StreamHandler()
        sh.setFormatter(formatter)
        logger.addHandler(sh)
        if log_path:
            try:
                fh = logging.FileHandler(log_path, encoding="utf-8")
                fh.setFormatter(formatter)
                logger.addHandler(fh)
            except OSError:
                pass


def create_app(db_path=None, log_path=None, api_key=None):
    app = Flask(__name__, static_folder=os.path.join(HERE, "static"))
    app.config["DB_PATH"] = db_path or os.environ.get("DB_PATH") or os.path.join(HERE, "monitor.db")
    app.config["LOG_PATH"] = log_path or os.environ.get("LOG_PATH") or os.path.join(HERE, "server.log")
    # When an API key is set, every POST must send it in the X-API-Key header.
    # Unset (local demo, tests) means writes are open, as before.
    app.config["API_KEY"] = api_key if api_key is not None else os.environ.get("DNHBM_API_KEY", "")
    setup_server_logging(app.config["LOG_PATH"])

    @app.before_request
    def require_api_key():
        key = app.config["API_KEY"]
        if key and request.method == "POST" and not hmac.compare_digest(
                request.headers.get("X-API-Key", ""), key):
            return jsonify({"error": "missing or wrong X-API-Key"}), 401

    def db():
        if "db" not in g:
            g.db = sqlite3.connect(app.config["DB_PATH"], timeout=30)
            g.db.row_factory = sqlite3.Row
            g.db.execute("PRAGMA foreign_keys = ON")
        return g.db

    @app.teardown_appcontext
    def close_db(_):
        conn = g.pop("db", None)
        if conn is not None:
            conn.close()

    with sqlite3.connect(app.config["DB_PATH"]) as init:
        init.execute("PRAGMA journal_mode = WAL")  # concurrent agents (TC13)
        init.executescript(SCHEMA)

    def err(msg, code=400):
        return jsonify({"error": msg}), code

    def is_num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool)

    # ---- FR-01: registration ----
    @app.post("/api/register")
    def register():
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            return err("JSON object required")
        node_id = data.get("node_id")
        if not isinstance(node_id, str) or not node_id.strip():
            return err("node_id (non-empty string) is required")
        interval = data.get("interval", DEFAULT_INTERVAL)
        if not is_num(interval) or interval <= 0:
            return err("interval must be a positive number")
        ts = now_iso()
        d = db()
        existing = d.execute("SELECT 1 FROM nodes WHERE node_id=?", (node_id,)).fetchone()
        if existing:  # TC02: update, never duplicate
            d.execute("UPDATE nodes SET hostname=?, os=?, interval_s=? WHERE node_id=?",
                      (data.get("hostname"), data.get("os"), interval, node_id))
        else:
            d.execute("INSERT INTO nodes VALUES (?,?,?,?,?,NULL)",
                      (node_id, data.get("hostname"), data.get("os"), interval, ts))
        d.commit()
        logger.info(f"Node {node_id} {'updated' if existing else 'registered'}: hostname={data.get('hostname')}")
        return jsonify({"status": "updated" if existing else "registered",
                        "node_id": node_id}), (200 if existing else 201)

    # ---- FR-05/06/07/08: ingestion (supports single object or batch array) ----
    @app.post("/api/metrics")
    def ingest():
        data = request.get_json(silent=True)
        if data is None or not isinstance(data, (dict, list)):
            return err("JSON object required")

        is_batch = isinstance(data, list)
        items = data if is_batch else [data]
        if not items:
            return err("JSON object required")

        d = db()
        stored_ts = None
        all_alerts = []

        for item in items:
            if not isinstance(item, dict):
                return err("JSON object required")
            item = {ALIASES.get(k, k): v for k, v in item.items()}
            node_id = item.get("node_id")
            if not isinstance(node_id, str) or not node_id.strip():
                return err("node_id (non-empty string) is required")
            values = {}
            for f in METRIC_FIELDS:
                v = item.get(f)
                if not is_num(v):
                    return err(f"{f} is required and must be a number")
                lo, hi = RANGES[f]
                if v < lo or (hi is not None and v > hi):
                    return err(f"{f} out of range")
                values[f] = v
            if not d.execute("SELECT 1 FROM nodes WHERE node_id=?", (node_id,)).fetchone():
                return err("unknown node; register first", 404)

            # Preserve client timestamp if present (e.g. from offline buffer)
            client_ts = item.get("timestamp")
            if isinstance(client_ts, str) and client_ts.strip():
                ts = client_ts.strip()
            else:
                ts = now_iso()
            stored_ts = ts

            d.execute("INSERT INTO metrics (node_id,timestamp,cpu_usage,memory_usage,"
                      "network_throughput,packet_loss,connection_count) VALUES (?,?,?,?,?,?,?)",
                      (node_id, ts, *[values[f] for f in METRIC_FIELDS]))
            d.execute("UPDATE nodes SET last_seen=? WHERE node_id=? AND (last_seen IS NULL OR last_seen <= ?)",
                      (ts, node_id, ts))

            for t in d.execute("SELECT * FROM thresholds").fetchall():
                val = values.get(t["metric_type"])
                if val is not None and val > t["limit_value"]:
                    msg = (f"{node_id}: {t['metric_type']} {val} exceeds "
                           f"threshold {t['limit_value']}")
                    d.execute("INSERT INTO alerts (node_id,threshold_id,message,timestamp) "
                              "VALUES (?,?,?,?)", (node_id, t["threshold_id"], msg, ts))
                    all_alerts.append(msg)
                    logger.warning(f"Threshold breach alert: {msg}")

        d.commit()
        if is_batch:
            return jsonify({"status": "stored", "count": len(items), "alerts": all_alerts}), 201
        return jsonify({"status": "stored", "timestamp": stored_ts, "alerts": all_alerts}), 201

    # ---- FR-09/10: metric queries ----
    @app.get("/api/metrics")
    def query_metrics():
        node_id = request.args.get("node_id")
        try:
            minutes = float(request.args.get("minutes", 60))
            limit = min(int(request.args.get("limit", 1000)), 10000)
        except ValueError:
            return err("minutes/limit must be numbers")
        since = datetime.fromtimestamp(time.time() - minutes * 60, timezone.utc).isoformat()
        sql, args = "SELECT * FROM metrics WHERE timestamp >= ?", [since]
        if node_id:
            sql += " AND node_id = ?"
            args.append(node_id)
        sql += " ORDER BY timestamp DESC LIMIT ?"
        args.append(limit)
        rows = [dict(r) for r in db().execute(sql, args)]
        return jsonify(list(reversed(rows)))

    # ---- FR-11: node status ----
    @app.get("/api/nodes")
    def nodes():
        out = []
        for r in db().execute("SELECT * FROM nodes ORDER BY node_id"):
            n = dict(r)
            if n["last_seen"]:
                try:
                    age = time.time() - datetime.fromisoformat(n["last_seen"]).timestamp()
                except Exception:
                    age = 0.0
                n["age_seconds"] = round(age, 1)
                n["status"] = "online" if age <= FRESH_INTERVALS * n["interval_s"] else "offline"
            else:
                n["age_seconds"], n["status"] = None, "registered"
            out.append(n)
        return jsonify(out)

    # ---- FR-12: alerts ----
    @app.get("/api/alerts")
    def alerts():
        # Alerts are never deleted by the server. ?minutes=30 returns only alerts from the
        # last 30 minutes; without it, the newest ?limit alerts are returned.
        limit = min(int(request.args.get("limit", 100)), 1000)
        sql, args = "SELECT * FROM alerts", []
        if request.args.get("minutes") is not None:
            try:
                minutes = float(request.args["minutes"])
            except ValueError:
                return err("minutes must be a number")
            since = datetime.fromtimestamp(time.time() - minutes * 60, timezone.utc)
            sql += " WHERE timestamp >= ?"
            args.append(since.isoformat(timespec="milliseconds"))
        sql += " ORDER BY alert_id DESC LIMIT ?"
        args.append(limit)
        return jsonify([dict(r) for r in db().execute(sql, args)])

    @app.get("/api/alerts/export")
    def export_alerts():
        # Full alert history as CSV, so it can be saved on your PC and looked at later.
        rows = db().execute("SELECT alert_id, node_id, threshold_id, timestamp, is_active, message "
                            "FROM alerts ORDER BY alert_id").fetchall()
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["alert_id", "node_id", "threshold_id", "timestamp", "is_active", "message"])
        w.writerows(tuple(r) for r in rows)
        return Response(buf.getvalue(), mimetype="text/csv",
                        headers={"Content-Disposition": "attachment; filename=alerts.csv"})

    @app.post("/api/alerts/clear")
    def clear_alerts():
        db().execute("UPDATE alerts SET is_active=0")
        db().commit()
        return jsonify({"status": "cleared"})

    # ---- FR-13: thresholds ----
    @app.get("/api/thresholds")
    def get_thresholds():
        return jsonify([dict(r) for r in db().execute("SELECT * FROM thresholds")])

    @app.post("/api/thresholds")
    def set_threshold():
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            return err("JSON object required")
        mt = ALIASES.get(data.get("metric_type"), data.get("metric_type"))
        limit = data.get("limit_value")
        if mt not in METRIC_FIELDS:
            return err(f"metric_type must be one of {METRIC_FIELDS}")
        if not is_num(limit) or limit < 0:
            return err("limit_value must be a non-negative number")
        d = db()
        d.execute("INSERT INTO thresholds (metric_type,limit_value) VALUES (?,?) "
                  "ON CONFLICT(metric_type) DO UPDATE SET limit_value=excluded.limit_value",
                  (mt, limit))
        d.commit()
        return jsonify({"status": "saved", "metric_type": mt, "limit_value": limit}), 200

    # ---- Offline Logs & Export ----
    @app.get("/api/logs")
    def get_logs():
        """Retrieve recent metric logs for offline inspection."""
        limit = min(int(request.args.get("limit", 100)), 1000)
        node_id = request.args.get("node_id")
        sql = "SELECT * FROM metrics"
        params = []
        if node_id:
            sql += " WHERE node_id = ?"
            params.append(node_id)
        sql += " ORDER BY timestamp DESC LIMIT ?"
        params.append(limit)
        metric_rows = [dict(r) for r in db().execute(sql, params)]
        return jsonify({"metrics": metric_rows})

    @app.get("/api/logs/export")
    def export_logs():
        """Export metrics logs in CSV or JSON format for offline viewing and analysis."""
        fmt = request.args.get("format", "csv").lower()
        node_id = request.args.get("node_id")
        limit = min(int(request.args.get("limit", 10000)), 50000)
        sql = ("SELECT timestamp, node_id, cpu_usage, memory_usage, "
               "network_throughput, packet_loss, connection_count FROM metrics")
        params = []
        if node_id:
            sql += " WHERE node_id = ?"
            params.append(node_id)
        sql += " ORDER BY timestamp DESC LIMIT ?"
        params.append(limit)
        rows = [dict(r) for r in db().execute(sql, params)]

        thresholds = {t["metric_type"]: t["limit_value"] for t in db().execute("SELECT metric_type, limit_value FROM thresholds")}

        if fmt == "json":
            return Response(
                json.dumps(rows, indent=2),
                mimetype="application/json",
                headers={"Content-Disposition": "attachment; filename=metrics_logs.json"}
            )

        if fmt in ("xlsx", "excel"):
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Telemetry Logs"
                headers = ["Timestamp", "Node ID", "CPU %", "Memory %", "Throughput (Mbps)", "Packet Loss %", "Connections"]
                ws.append(headers)

                header_fill = PatternFill(start_color="F6F8FA", end_color="F6F8FA", fill_type="solid")
                header_font = Font(bold=True, color="1F2D3D")
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.fill = header_fill
                    cell.font = header_font

                breach_fill = PatternFill(start_color="FFEBE9", end_color="FFEBE9", fill_type="solid")
                breach_font = Font(bold=True, color="CF222E")

                for r_idx, r in enumerate(rows, start=2):
                    ws.append([r["timestamp"], r["node_id"], r["cpu_usage"], r["memory_usage"],
                               r["network_throughput"], r["packet_loss"], r["connection_count"]])
                    for c_offset, f in enumerate(METRIC_FIELDS):
                        val = r[f]
                        lim = thresholds.get(f)
                        if lim is not None and val is not None and val > lim:
                            target_cell = ws.cell(row=r_idx, column=3 + c_offset)
                            target_cell.fill = breach_fill
                            target_cell.font = breach_font

                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                return Response(
                    buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=metrics_logs.xlsx"}
                )
            except ImportError:
                pass  # Fallback to CSV if openpyxl is unavailable

        if fmt == "html":
            html_rows = []
            for r in rows:
                cells = [f"<td>{r['timestamp']}</td>", f"<td>{r['node_id']}</td>"]
                for f in METRIC_FIELDS:
                    val = r[f]
                    lim = thresholds.get(f)
                    if lim is not None and val is not None and val > lim:
                        cells.append(f'<td style="color:#cf222e;font-weight:bold;background:#ffebe9;">{val}</td>')
                    else:
                        cells.append(f"<td>{val}</td>")
                html_rows.append(f"<tr>{''.join(cells)}</tr>")

            html = f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>Exported Telemetry Logs</title>
<style>
  body {{ font-family: system-ui, sans-serif; margin: 20px; background: #fafafa; color: #222; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; background: #fff; }}
  th, td {{ border: 1px solid #e1e4e8; padding: 6px 10px; text-align: left; }}
  th {{ background: #f6f8fa; font-weight: 600; }}
</style>
</head>
<body>
<h2>Exported Telemetry Logs</h2>
<table>
<thead>
  <tr><th>Timestamp</th><th>Node</th><th>CPU %</th><th>Memory %</th><th>Throughput (Mbps)</th><th>Packet Loss %</th><th>Connections</th></tr>
</thead>
<tbody>
  {''.join(html_rows)}
</tbody>
</table>
</body>
</html>"""
            return Response(
                html,
                mimetype="text/html",
                headers={"Content-Disposition": "attachment; filename=metrics_logs.html"}
            )

        # Default: CSV format with breached field marked
        out = io.StringIO()
        fieldnames = ["timestamp", "node_id", "cpu_usage", "memory_usage",
                      "network_throughput", "packet_loss", "connection_count"]
        writer = csv.DictWriter(out, fieldnames=fieldnames)
        writer.writeheader()
        for r in rows:
            row_dict = dict(r)
            for f in METRIC_FIELDS:
                val = r[f]
                lim = thresholds.get(f)
                if lim is not None and val is not None and val > lim:
                    row_dict[f] = f"{val} [ALERT]"
            writer.writerow(row_dict)
        return Response(
            out.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": "attachment; filename=metrics_logs.csv"}
        )

    @app.get("/dashboar")
    @app.get("/dashboard")
    def dashboard():
        return send_from_directory(os.path.join(HERE, "static"), "dashboard.html")

    @app.get("/")
    def index():
        return dashboard()

    return app


if __name__ == "__main__":
    app = create_app()
    host = os.environ.get("HOST", os.environ.get("NETMON_HOST", "127.0.0.1"))
    port = int(os.environ.get("PORT", os.environ.get("NETMON_PORT", 5000)))
    if port == 5000 and host == "127.0.0.1":
        import threading
        threading.Thread(target=lambda: app.run(host=host, port=5001, threaded=True), daemon=True).start()
    app.run(host=host, port=port, threaded=True)
